# -*- coding: utf-8 -*-
"""Construcción del .docx a partir de los bloques de contenido."""

from __future__ import annotations

import re
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsmap, qn
    from docx.shared import Cm, Pt
except ImportError:  # pragma: no cover
    raise SystemExit(
        "Falta python-docx.\n"
        "Corra el generador con el python del entorno del proyecto:\n"
        "    .venv/bin/python generar_boletin.py\n"
        "o cree el entorno si no existe:\n"
        "    python3 -m venv .venv\n"
        "    .venv/bin/pip install -r requirements.txt"
    )

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


# ---------------------------------------------------------------------------
# Utilidades de bajo nivel
# ---------------------------------------------------------------------------

def _borde(elemento, lado, tamano, color="000000"):
    """Agrega un borde a un <w:tcPr> (estilo booktabs)."""
    bordes = elemento.find(qn("w:tcBorders"))
    if bordes is None:
        bordes = OxmlElement("w:tcBorders")
        elemento.append(bordes)
    linea = OxmlElement(f"w:{lado}")
    if tamano == 0:
        linea.set(qn("w:val"), "none")
        linea.set(qn("w:sz"), "0")
    else:
        linea.set(qn("w:val"), "single")
        linea.set(qn("w:sz"), str(tamano))
    linea.set(qn("w:space"), "0")
    linea.set(qn("w:color"), color)
    bordes.append(linea)


def _propiedades_celda(celda):
    tcPr = celda._tc.get_or_add_tcPr()
    return tcPr


def _fijar_fuente(run, cfg, tamano=None, negrita=False, cursiva=False):
    run.font.name = cfg["fuente"]
    run.font.size = Pt(tamano or cfg["tamano_pt"])
    run.bold = negrita
    run.italic = cursiva
    # Word necesita el nombre repetido para los alfabetos no latinos.
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(atributo), cfg["fuente"])


def _escribir(parrafo, texto, cfg, tamano=None, negrita=False, cursiva=False):
    """Escribe el texto interpretando **negrita** y *cursiva*."""
    for trozo in INLINE.split(texto):
        if not trozo:
            continue
        if trozo.startswith("**") and trozo.endswith("**") and len(trozo) > 4:
            run = parrafo.add_run(trozo[2:-2])
            _fijar_fuente(run, cfg, tamano, True, cursiva)
        elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
            run = parrafo.add_run(trozo[1:-1])
            _fijar_fuente(run, cfg, tamano, negrita, True)
        else:
            run = parrafo.add_run(trozo)
            _fijar_fuente(run, cfg, tamano, negrita, cursiva)
    return parrafo


# ---------------------------------------------------------------------------
# Constructor del documento
# ---------------------------------------------------------------------------

class Documento:

    def __init__(self, cfg):
        self.cfg = cfg
        self.doc = Document()
        self._configurar()

    # -- configuración general --------------------------------------------
    def _configurar(self):
        cfg = self.cfg
        for seccion in self.doc.sections:
            seccion.page_width = Cm(cfg["ancho_pagina_cm"])
            seccion.page_height = Cm(cfg["alto_pagina_cm"])
            seccion.left_margin = Cm(cfg["margen_lateral_cm"])
            seccion.right_margin = Cm(cfg["margen_lateral_cm"])
            seccion.top_margin = Cm(cfg["margen_vertical_cm"])
            seccion.bottom_margin = Cm(cfg["margen_vertical_cm"])

        normal = self.doc.styles["Normal"]
        normal.font.name = cfg["fuente"]
        normal.font.size = Pt(cfg["tamano_pt"])
        rPr = normal.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(atributo), cfg["fuente"])
        formato = normal.paragraph_format
        formato.line_spacing = cfg["interlineado"]
        formato.space_after = Pt(cfg["espacio_despues_pt"])

    # -- bloques -----------------------------------------------------------
    def parrafo(self, texto, tamano=None, negrita=False, alineacion=None,
                estilo=None, espacio_despues=None, interlineado=None):
        p = self.doc.add_paragraph(style=estilo)
        if alineacion is not None:
            p.paragraph_format.alignment = alineacion
        if espacio_despues is not None:
            p.paragraph_format.space_after = Pt(espacio_despues)
        if interlineado is not None:
            p.paragraph_format.line_spacing = interlineado
        _escribir(p, texto, self.cfg, tamano, negrita)
        return p

    def titulo(self, texto, nota=None):
        p = self.parrafo(
            texto,
            tamano=self.cfg["tamano_pt"] + 4,
            negrita=True,
            alineacion=WD_ALIGN_PARAGRAPH.CENTER,
        )
        if nota:
            self.nota_al_pie(p, nota)
        return p

    # -- Notas al pie ----------------------------------------------------
    # python-docx no las soporta, asi que se arma la parte word/footnotes.xml
    # a mano y se le cuelga la referencia al parrafo.
    def nota_al_pie(self, parrafo, texto, tamano=None):
        indice = self._asegurar_footnotes()

        cuerpo = OxmlElement("w:footnote")
        cuerpo.set(qn("w:id"), str(indice))
        pf = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "both")
        ppr.append(jc)
        pf.append(ppr)

        # marca del numero dentro de la nota
        r_marca = OxmlElement("w:r")
        rpr_m = OxmlElement("w:rPr")
        va = OxmlElement("w:vertAlign"); va.set(qn("w:val"), "superscript")
        sz_m = OxmlElement("w:sz"); sz_m.set(qn("w:val"), str(int((tamano or self.cfg["tamano_nota_pt"]) * 2)))
        rpr_m.append(va); rpr_m.append(sz_m)
        r_marca.append(rpr_m)
        r_marca.append(OxmlElement("w:footnoteRef"))
        pf.append(r_marca)

        r_txt = OxmlElement("w:r")
        rpr_t = OxmlElement("w:rPr")
        sz_t = OxmlElement("w:sz"); sz_t.set(qn("w:val"), str(int((tamano or self.cfg["tamano_nota_pt"]) * 2)))
        rf = OxmlElement("w:rFonts")
        rf.set(qn("w:ascii"), self.cfg["fuente"]); rf.set(qn("w:hAnsi"), self.cfg["fuente"])
        rpr_t.append(rf); rpr_t.append(sz_t)
        r_txt.append(rpr_t)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = " " + texto
        r_txt.append(t)
        pf.append(r_txt)

        cuerpo.append(pf)
        self._footnotes.append(cuerpo)

        # referencia en el texto
        run = parrafo.add_run()
        rpr = run._r.get_or_add_rPr()
        va2 = OxmlElement("w:vertAlign"); va2.set(qn("w:val"), "superscript")
        rpr.append(va2)
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), str(indice))
        run._r.append(ref)
        return indice

    def _asegurar_footnotes(self):
        """Crea word/footnotes.xml la primera vez y devuelve el proximo id."""
        if getattr(self, "_footnotes", None) is not None:
            self._siguiente_nota += 1
            return self._siguiente_nota

        ns = " ".join(f'xmlns:{k}="{v}"' for k, v in nsmap.items())
        base = (
            f'<w:footnotes {ns}>'
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr>'
            '<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
            '<w:r><w:separator/></w:r></w:p></w:footnote>'
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr>'
            '<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
            '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
            '</w:footnotes>'
        )
        self._footnotes = parse_xml(base)
        self._siguiente_nota = 1

        parte = Part(
            PackURI("/word/footnotes.xml"),
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.footnotes+xml",
            b"",
            self.doc.part.package,
        )
        self._parte_footnotes = parte
        self.doc.part.relate_to(parte, RT.FOOTNOTES)
        return 1

    def seccion(self, texto):
        p = self.parrafo(
            texto, tamano=self.cfg["tamano_pt"] + 2, negrita=True
        )
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.keep_with_next = True
        return p

    def subseccion(self, texto):
        p = self.parrafo(texto, negrita=True)
        p.paragraph_format.keep_with_next = True
        return p

    def vineta(self, texto):
        p = self.doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = self.cfg["interlineado"]
        p.paragraph_format.space_after = Pt(self.cfg["espacio_despues_pt"])
        _escribir(p, texto, self.cfg)
        return p

    def salto_de_pagina(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -- figuras -----------------------------------------------------------
    def encabezado_figura(self, texto):
        return self.parrafo(
            texto,
            negrita=False,
            alineacion=WD_ALIGN_PARAGRAPH.CENTER,
            espacio_despues=2,
            interlineado=1.0,
        )

    def pie(self, texto):
        return self.parrafo(
            texto,
            tamano=self.cfg["tamano_nota_pt"],
            espacio_despues=0,
            interlineado=1.0,
        )

    def imagen(self, ruta):
        p = self.doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(ruta), width=Cm(self.cfg["ancho_imagen_cm"]))
        return p

    def marcador_faltante(self, descripcion):
        """Recuadro visible cuando un gráfico todavía no fue generado."""
        tabla = self.doc.add_table(rows=1, cols=1)
        tabla.style = "Table Grid"
        celda = tabla.cell(0, 0)
        celda.text = ""
        p = celda.paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _escribir(
            p,
            f"[FALTA LA IMAGEN] {descripcion}",
            self.cfg,
            tamano=self.cfg["tamano_nota_pt"],
            negrita=True,
        )
        self.doc.add_paragraph()
        return tabla

    # -- tablas ------------------------------------------------------------
    def tabla(self, encabezados, filas):
        cfg = self.cfg
        tabla = self.doc.add_table(rows=1, cols=len(encabezados))
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabla.autofit = True

        def _celda(celda, texto, negrita=False, primera_columna=False):
            celda.text = ""
            p = celda.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if not primera_columna:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _escribir(p, texto, cfg, tamano=cfg["tamano_tabla_pt"], negrita=negrita)

        for i, texto in enumerate(encabezados):
            celda = tabla.rows[0].cells[i]
            _celda(celda, texto, negrita=True, primera_columna=(i == 0))
            tcPr = _propiedades_celda(celda)
            _borde(tcPr, "top", 12)
            _borde(tcPr, "bottom", 6)
            _borde(tcPr, "left", 0)
            _borde(tcPr, "right", 0)

        for numero, fila in enumerate(filas):
            celdas = tabla.add_row().cells
            ultima = numero == len(filas) - 1
            for i, texto in enumerate(fila):
                _celda(celdas[i], texto, primera_columna=(i == 0))
                tcPr = _propiedades_celda(celdas[i])
                _borde(tcPr, "left", 0)
                _borde(tcPr, "right", 0)
                _borde(tcPr, "top", 0)
                _borde(tcPr, "bottom", 12 if ultima else 0)
        return tabla

    def cuadro(self, titulo, parrafos):
        if titulo:
            self.parrafo(
                titulo,
                negrita=True,
                alineacion=WD_ALIGN_PARAGRAPH.CENTER,
                espacio_despues=2,
                interlineado=1.0,
            )
        tabla = self.doc.add_table(rows=1, cols=1)
        tabla.style = "Table Grid"
        celda = tabla.cell(0, 0)
        celda.text = ""
        primera = True
        for texto in parrafos:
            p = celda.paragraphs[0] if primera else celda.add_paragraph()
            primera = False
            p.paragraph_format.line_spacing = self.cfg["interlineado"]
            p.paragraph_format.space_after = Pt(self.cfg["espacio_despues_pt"])
            _escribir(p, texto, self.cfg, tamano=self.cfg["tamano_nota_pt"] + 1)
        self.doc.add_paragraph()
        return tabla

    def equipo(self, integrantes):
        tabla = self.doc.add_table(rows=0, cols=1)
        tabla.style = "Table Grid"
        for cargo, personas in integrantes:
            fila = tabla.add_row().cells[0]
            fila.text = ""
            _escribir(fila.paragraphs[0], cargo, self.cfg, negrita=True)
            fila.paragraphs[0].paragraph_format.space_after = Pt(0)
            for persona in personas:
                p = fila.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                _escribir(p, persona, self.cfg)
        self.doc.add_paragraph()
        return tabla

    # -- guardado ----------------------------------------------------------
    def guardar(self, ruta: Path):
        ruta = Path(ruta)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        if getattr(self, "_footnotes", None) is not None:
            from lxml import etree
            self._parte_footnotes._blob = etree.tostring(
                self._footnotes, xml_declaration=True,
                encoding="UTF-8", standalone=True)
        self.doc.save(str(ruta))
        return ruta
